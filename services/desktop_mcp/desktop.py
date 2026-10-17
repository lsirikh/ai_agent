"""
데스크톱 작업: 정해진 프로그램 실행 · 문서 작성/저장 · 저장한 파일 확인.

EdgeVoice 의 로컬 LLM 이 MCP(server.py)로 부른다. 이 PC(Windows) 화면에서 돈다.
- 프로그램은 APPS 에 있는 것만 실행한다.
- 파일은 Downloads/Tests 안에만 저장하고, 같은 이름이 있으면 덮어쓰지 않는다.
- 파일은 라이브러리(python-docx·openpyxl·python-pptx)로 만들어 저장한 뒤
  그 프로그램으로 열어 화면에 보여 준다. Office COM 자동화는 이 PC 의
  'Office 인증 마법사' 창에 막혀 엑셀은 저장되지 않고 파워포인트는 멈췄다(실측).
"""

import os
import re
from pathlib import Path
from typing import Callable, Optional

DEFAULT_FOLDER = Path.home() / "Downloads" / "Tests"
MAX_CONTENT_CHARS = 20000

# 이름 → (실행 파일, 저장 확장자)
APPS = {
    "notepad": ("notepad.exe", ".txt"),
    "word": ("winword.exe", ".docx"),
    "excel": ("excel.exe", ".xlsx"),
    "powerpoint": ("powerpnt.exe", ".pptx"),
    "calculator": ("calc.exe", None),
    "paint": ("mspaint.exe", None),
}
_ALIASES = {
    "메모장": "notepad",
    "노트패드": "notepad",
    "텍스트": "notepad",
    "워드": "word",
    "한글워드": "word",
    "ms워드": "word",
    "엑셀": "excel",
    "파워포인트": "powerpoint",
    "피피티": "powerpoint",
    "ppt": "powerpoint",
    "계산기": "calculator",
    "calc": "calculator",
    "그림판": "paint",
    "mspaint": "paint",
}
KOREAN_NAMES = {
    "notepad": "메모장",
    "word": "워드",
    "excel": "엑셀",
    "powerpoint": "파워포인트",
    "calculator": "계산기",
    "paint": "그림판",
}
_UNSAFE_NAME = re.compile(r'[\\/:*?"<>|\x00-\x1f]|\.\.')


def normalize_app(name: str) -> str:
    key = (name or "").strip().lower().replace(" ", "")
    key = _ALIASES.get(key, key)
    if key not in APPS:
        allowed = ", ".join(KOREAN_NAMES.values())
        raise ValueError(f"실행할 수 없는 프로그램: {name} (가능: {allowed})")
    return key


def target_path(folder: Path, filename: str, app: str) -> Path:
    """folder 안의 저장 경로. 경로 문자는 지우고, 있으면 ' (2)' 처럼 새 이름."""
    suffix = APPS[app][1]
    stem = Path(_UNSAFE_NAME.sub("", filename or "").strip()).stem.strip(". ") or "문서"
    stem = stem[:80]
    path = folder / f"{stem}{suffix}"
    n = 2
    while path.exists():
        path = folder / f"{stem} ({n}){suffix}"
        n += 1
    return path


def _cell(value: str):
    value = value.strip()
    try:
        return int(value)
    except ValueError:
        try:
            return float(value)
        except ValueError:
            return value


def table_rows(content: str) -> list:
    """줄은 행, 탭·쉼표·| 는 칸."""
    rows = []
    for line in content.splitlines():
        if line.strip():
            rows.append([_cell(c) for c in re.split(r"\t|,|\|", line)])
    return rows


def _launch(path_or_exe) -> None:
    os.startfile(str(path_or_exe))  # noqa: S606 — 허용 목록 안의 프로그램·저장한 파일만


def open_app(name: str, launch: Callable = _launch) -> str:
    app = normalize_app(name)
    launch(APPS[app][0])
    return app


def _save_word(path: Path, content: str) -> None:
    import docx

    document = docx.Document()
    for line in content.split("\n"):
        document.add_paragraph(line)
    document.save(str(path))


def _save_excel(path: Path, content: str) -> None:
    import openpyxl

    book = openpyxl.Workbook()
    for row in table_rows(content):
        book.active.append(row)
    book.save(str(path))


_NUMBERED_HEADING = re.compile(r"^\d+[.)]\s+")


def slide_blocks(content: str) -> list:
    """슬라이드마다 [제목, 본문 줄…]. 빈 줄로 나누고, 빈 줄이 없으면
    맨 앞의 번호 줄('1. ', '2) ')에서 나눈다(로컬 LLM 이 그렇게 보냈다 — 실측)."""
    text = content.strip()
    blocks = [b for b in re.split(r"\n\s*\n", text) if b.strip()]
    if len(blocks) <= 1:
        blocks = re.split(r"\n(?=\d+[.)]\s)", text)
    slides = []
    for block in blocks:
        lines = [line.strip() for line in block.strip().split("\n") if line.strip()]
        if lines:
            lines[0] = _NUMBERED_HEADING.sub("", lines[0])
            slides.append(lines)
    return slides or [[""]]


def _save_powerpoint(path: Path, content: str) -> None:
    """슬라이드마다 첫 줄은 제목, 나머지는 본문(slide_blocks)."""
    import pptx

    deck = pptx.Presentation()
    for lines in slide_blocks(content):
        slide = deck.slides.add_slide(deck.slide_layouts[1])  # 제목과 내용
        slide.shapes.title.text = lines[0]
        slide.placeholders[1].text = "\n".join(lines[1:])
    deck.save(str(path))


_SAVERS = {"word": _save_word, "excel": _save_excel, "powerpoint": _save_powerpoint}


def write_document(
    app_name: str,
    filename: str,
    content: str,
    folder: Path = DEFAULT_FOLDER,
    launch: Callable = _launch,
) -> Path:
    app = normalize_app(app_name)
    if APPS[app][1] is None:
        raise ValueError(f"{KOREAN_NAMES[app]}은(는) 문서를 저장할 수 없습니다")
    content = (content or "")[:MAX_CONTENT_CHARS]
    folder.mkdir(parents=True, exist_ok=True)
    path = target_path(folder, filename, app)
    if app == "notepad":
        path.write_text(content, encoding="utf-8")
    else:
        _SAVERS[app](path, content)
    launch(path)
    return path


def read_document(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".txt":
        return path.read_text(encoding="utf-8")
    if suffix == ".docx":
        import docx

        return "\n".join(p.text for p in docx.Document(str(path)).paragraphs)
    if suffix == ".xlsx":
        import openpyxl

        sheet = openpyxl.load_workbook(str(path), read_only=True).active
        return "\n".join(
            ", ".join("" if v is None else str(v) for v in row)
            for row in sheet.iter_rows(values_only=True)
        )
    if suffix == ".pptx":
        import pptx

        slides = []
        for slide in pptx.Presentation(str(path)).slides:
            texts = [s.text_frame.text for s in slide.shapes if s.has_text_frame]
            slides.append("\n".join(t for t in texts if t))
        return "\n\n".join(slides)
    raise ValueError(f"읽을 수 없는 형식: {path.suffix}")


def find_file(folder: Path, filename: str) -> Optional[Path]:
    name = Path(_UNSAFE_NAME.sub("", filename or "")).name
    path = folder / name
    return path if name and path.is_file() else None


def list_files(folder: Path = DEFAULT_FOLDER) -> str:
    if not folder.is_dir():
        return f"{folder} 폴더가 아직 없습니다"
    files = sorted(folder.iterdir(), key=lambda p: p.stat().st_mtime, reverse=True)
    lines = [f"{p.name} ({p.stat().st_size:,}바이트)" for p in files if p.is_file()]
    return "\n".join(lines) or "저장된 파일이 없습니다"
